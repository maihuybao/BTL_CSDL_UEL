"""DbFormController: lớp CRUD chung nối bảng <-> form theo col_mappings.

Tách nguyên văn từ Main_Controller.py — không đổi logic.
"""
import sqlite3

from PyQt6 import QtWidgets, QtGui, QtCore
from PyQt6.QtWidgets import QTableWidgetItem, QMessageBox

from app.config import DB_PATH, HAS_QTA, current_dir
if HAS_QTA:
    from app.config import qta
from app.helpers import get_widget_value, set_widget_value
from app.db_logic import apply_order_status_effects, recalc_order_total


class DbFormController(QtCore.QObject):
    def __init__(self, main_controller, ui_page, table_widget, db_table, pk_col, col_mappings, refresh_callback):
        super().__init__()
        self.main_controller = main_controller
        self.ui = ui_page
        self.table_widget = table_widget
        self.db_table = db_table
        self.pk_col = pk_col
        self.col_mappings = col_mappings  # list of tuples: (db_col, widget, widget_type)
        self.refresh_callback = refresh_callback
        self.current_action = None
        self.selected_pk = None

        self.setup_connections()
        self.set_form_enabled(False)

    def setup_connections(self):
        # Bắt buộc chọn theo dòng để click vào ô bất kỳ cũng đổ được dữ liệu lên form
        self.table_widget.setSelectionBehavior(QtWidgets.QAbstractItemView.SelectionBehavior.SelectRows)
        self.table_widget.setSelectionMode(QtWidgets.QAbstractItemView.SelectionMode.SingleSelection)
        self.table_widget.setEditTriggers(QtWidgets.QAbstractItemView.EditTrigger.NoEditTriggers)
        self.table_widget.itemSelectionChanged.connect(self.sync_table_to_form)
        self.table_widget.cellClicked.connect(lambda r, c: self.sync_table_to_form())
        if hasattr(self.ui, 'btnThem'):
            self.ui.btnThem.clicked.connect(self.action_them)
        if hasattr(self.ui, 'btnSua'):
            self.ui.btnSua.clicked.connect(self._on_edit_button)
        if hasattr(self.ui, 'btnXoa'):
            self.ui.btnXoa.clicked.connect(self.action_xoa)
        if hasattr(self.ui, 'btnLuu'):
            self.ui.btnLuu.clicked.connect(self.action_luu)
        if hasattr(self.ui, 'btnHuy'):
            self.ui.btnHuy.clicked.connect(self.action_huy)
        if hasattr(self.ui, 'btnRefresh'):
            self.ui.btnRefresh.clicked.connect(self.action_refresh)
        if hasattr(self.ui, 'btnSearch'):
            self.ui.btnSearch.clicked.connect(self.action_search)
        if hasattr(self.ui, 'btnXuatExcel'):
            self.ui.btnXuatExcel.clicked.connect(self.action_xuat_excel)

    def set_form_enabled(self, enabled=True):
        for db_col, widget, widget_type in self.col_mappings:
            if widget:
                widget.setEnabled(enabled)
        
        if hasattr(self.ui, 'btnThem'):
            self.ui.btnThem.setEnabled(not enabled)
        if hasattr(self.ui, 'btnXoa'):
            self.ui.btnXoa.setEnabled(not enabled)
        # Nút Sửa toggle thành Lưu: đang sửa/thêm -> "Lưu", ngược lại -> "Sửa"
        if hasattr(self.ui, 'btnSua'):
            self.ui.btnSua.setEnabled(True)
            if enabled:
                self.ui.btnSua.setText("Lưu")
                if HAS_QTA:
                    self.ui.btnSua.setIcon(qta.icon('fa5s.save', color='white'))
            else:
                self.ui.btnSua.setText("Sửa")
                if HAS_QTA:
                    self.ui.btnSua.setIcon(qta.icon('fa5s.edit', color='white'))
        # Ẩn nút Lưu riêng — mọi submit đi qua nút Sửa/Lưu đã toggle
        if hasattr(self.ui, 'btnLuu'):
            self.ui.btnLuu.setVisible(False)
        if hasattr(self.ui, 'btnHuy'):
            self.ui.btnHuy.setEnabled(enabled)

    def sync_table_to_form(self):
        if self.current_action is not None:
            return
        
        selected_items = self.table_widget.selectedItems()
        if not selected_items:
            self.selected_pk = None
            return
            
        row = selected_items[0].row()
        self.selected_pk = self.table_widget.item(row, 0).text()
        
        try:
            conn = sqlite3.connect(DB_PATH)
            cursor = conn.cursor()
            
            if self.db_table == "LIVESTREAM_SAN_PHAM":
                parts = self.selected_pk.split("-")
                if len(parts) >= 2:
                    cursor.execute("""
                        SELECT lsp.MaLive, lsp.MaSP, p.SoLuongTon, p.GiaBan
                        FROM LIVESTREAM_SAN_PHAM lsp
                        JOIN SAN_PHAM p ON lsp.MaSP = p.MaSP
                        WHERE lsp.MaLive = ? AND lsp.MaSP = ?
                    """, (parts[0], parts[1]))
                    row_data = cursor.fetchone()
                else:
                    row_data = None
            elif self.db_table == "CHI_TIET_DON_HANG":
                parts = self.selected_pk.split("-")
                if len(parts) >= 2:
                    cursor.execute("SELECT MaDonHang, MaSP, SoLuong, DonGia, ThanhTien FROM CHI_TIET_DON_HANG WHERE MaDonHang = ? AND MaSP = ?", (parts[0], parts[1]))
                    row_data = cursor.fetchone()
                else:
                    row_data = None
            elif self.db_table == "DON_HANG":
                cursor.execute("""
                    SELECT 
                        dh.NgayDat,
                        bl.MaLive,
                        kh.HoTen,
                        kh.SoDienThoai,
                        kh.DiaChi,
                        dh.MaVoucher,
                        dh.TrangThaiDH,
                        dh.TongTien
                    FROM DON_HANG dh
                    LEFT JOIN KHACH_HANG kh ON dh.MaKhachHang = kh.MaKhachHang
                    LEFT JOIN BINH_LUAN bl ON dh.MaBinhLuan = bl.MaBinhLuan
                    WHERE dh.MaDonHang = ?
                """, (self.selected_pk,))
                order_res = cursor.fetchone()
                if order_res:
                    row_data = (order_res[0], order_res[1], order_res[2], order_res[3], order_res[4], order_res[5], order_res[6], order_res[7])
                else:
                    row_data = None
                
                # Fetch order products (kèm ảnh sản phẩm nếu có trong thư mục dự án)
                cursor.execute("""
                    SELECT p.TenSP, ct.SoLuong, ct.DonGia, p.HinhAnh
                    FROM CHI_TIET_DON_HANG ct
                    JOIN SAN_PHAM p ON ct.MaSP = p.MaSP
                    WHERE ct.MaDonHang = ?
                """, (self.selected_pk,))
                products = cursor.fetchall()
                if hasattr(self.ui, "tblOrderProductsList"):
                    self.main_controller.populate_table(self.ui.tblOrderProductsList,
                                                        [p[:3] for p in products])
                    # Gắn icon ảnh vào cột tên sản phẩm
                    for r, prod in enumerate(products):
                        img = prod[3]
                        item = self.ui.tblOrderProductsList.item(r, 0)
                        if not (img and item):
                            continue
                        for candidate in (os.path.join(current_dir, img),
                                          os.path.join(current_dir, os.path.splitext(img)[0] + ".png")):
                            if os.path.exists(candidate):
                                item.setIcon(QtGui.QIcon(candidate))
                                self.ui.tblOrderProductsList.setIconSize(QtCore.QSize(32, 32))
                                break
            else:
                cols = [col for col, _, _ in self.col_mappings if col is not None]
                query = f"SELECT {', '.join(cols)} FROM {self.db_table} WHERE {self.pk_col} = ?"
                cursor.execute(query, (self.selected_pk,))
                row_data = cursor.fetchone()
            conn.close()
            
            if row_data:
                for idx, (db_col, widget, widget_type) in enumerate(self.col_mappings):
                    if widget:
                        set_widget_value(widget, widget_type, row_data[idx])
        except Exception as e:
            print(f"Error syncing {self.db_table} to form: {e}")

    def get_pk_widget(self):
        for db_col, widget, widget_type in self.col_mappings:
            if db_col == self.pk_col:
                return widget
        return None

    def get_pk_widget_type(self):
        for db_col, widget, widget_type in self.col_mappings:
            if db_col == self.pk_col:
                return widget_type
        return "text"

    def clear_form(self):
        for db_col, widget, widget_type in self.col_mappings:
            if widget:
                if widget_type == "text":
                    widget.clear()
                elif widget_type == "combo_text":
                    widget.setCurrentIndex(0)
                elif widget_type == "combo_index":
                    widget.setCurrentIndex(0)
                elif widget_type == "spin":
                    widget.setValue(0)
                elif widget_type == "date":
                    widget.setDate(QtCore.QDate.currentDate())
                elif widget_type == "datetime":
                    widget.setDateTime(QtCore.QDateTime.currentDateTime())

    def generate_new_pk(self):
        prefix = ""
        if self.db_table == "KHACH_HANG": prefix = "KH"
        elif self.db_table == "SAN_PHAM": prefix = "SP"
        elif self.db_table == "NGUOI_BAN": prefix = "NB"
        elif self.db_table == "LIVESTREAM": prefix = "LIVE"
        elif self.db_table == "DON_HANG": prefix = "DH"
        elif self.db_table == "BINH_LUAN": prefix = "BL"
        elif self.db_table == "VOUCHER": prefix = "VOUCHER"
        elif self.db_table == "HOA_DON": prefix = "HD"
        return self.generate_new_pk_generic(self.db_table, self.pk_col, prefix)

    def generate_new_pk_generic(self, db_table, pk_col, prefix):
        try:
            conn = sqlite3.connect(DB_PATH)
            cursor = conn.cursor()
            cursor.execute(f"SELECT {pk_col} FROM {db_table}")
            pks = [row[0] for row in cursor.fetchall() if row[0] and row[0].startswith(prefix)]
            conn.close()
            
            max_num = 0
            for pk in pks:
                num_part = pk[len(prefix):]
                if num_part.isdigit():
                    max_num = max(max_num, int(num_part))
            
            return f"{prefix}{max_num + 1:02d}"
        except Exception as e:
            print(f"Error generating PK: {e}")
            return prefix + "99"

    def _on_edit_button(self):
        # Nút Sửa/Lưu gộp: đang mở form (THEM/SUA) -> submit; ngược lại -> mở sửa
        if self.current_action in ("THEM", "SUA"):
            self.action_luu()
        else:
            self.action_sua()

    def action_them(self):
        self.current_action = "THEM"
        self.selected_pk = None
        self.clear_form()
        self.set_form_enabled(True)
        pk_widget = self.get_pk_widget()
        if pk_widget:
            pk_widget.setEnabled(True)

    def action_sua(self):
        selected_items = self.table_widget.selectedItems()
        if not selected_items:
            QMessageBox.warning(self.main_controller, "Cảnh báo", "Vui lòng chọn một bản ghi trên bảng để sửa!")
            return
        self.current_action = "SUA"
        self.set_form_enabled(True)
        pk_widget = self.get_pk_widget()
        if pk_widget:
            pk_widget.setEnabled(False)

    def action_xoa(self):
        selected_items = self.table_widget.selectedItems()
        if not selected_items:
            QMessageBox.warning(self.main_controller, "Cảnh báo", "Vui lòng chọn một bản ghi trên bảng để xóa!")
            return
            
        row = selected_items[0].row()
        pk_val = self.table_widget.item(row, 0).text()
        
        confirm = QMessageBox.question(
            self.main_controller, "Xác nhận xóa",
            f"Bạn có chắc chắn muốn xóa bản ghi {pk_val} không?",
            QMessageBox.StandardButton.Yes | QMessageBox.StandardButton.No
        )
        
        if confirm == QMessageBox.StandardButton.Yes:
            try:
                conn = sqlite3.connect(DB_PATH)
                cursor = conn.cursor()
                
                if self.db_table == "LIVESTREAM_SAN_PHAM":
                    parts = pk_val.split("-")
                    cursor.execute("DELETE FROM LIVESTREAM_SAN_PHAM WHERE MaLive = ? AND MaSP = ?", (parts[0], parts[1]))
                elif self.db_table == "CHI_TIET_DON_HANG":
                    parts = pk_val.split("-")
                    cursor.execute("DELETE FROM CHI_TIET_DON_HANG WHERE MaDonHang = ? AND MaSP = ?", (parts[0], parts[1]))

                    # Update Totals (đã trừ voucher nếu có)
                    recalc_order_total(cursor, parts[0])
                elif self.db_table == "DON_HANG":
                    # Delete invoices and order details first to avoid constraint violation
                    cursor.execute("DELETE FROM HOA_DON WHERE MaDonHang = ?", (pk_val,))
                    cursor.execute("DELETE FROM CHI_TIET_DON_HANG WHERE MaDonHang = ?", (pk_val,))
                    cursor.execute(f"DELETE FROM {self.db_table} WHERE {self.pk_col} = ?", (pk_val,))
                else:
                    cursor.execute(f"DELETE FROM {self.db_table} WHERE {self.pk_col} = ?", (pk_val,))
                
                conn.commit()
                conn.close()
                self.refresh_callback()
                self.clear_form()
                # Đồng bộ các màn hình liên quan (bảng, combobox, thẻ số liệu, biểu đồ)
                if hasattr(self.main_controller, 'refresh_dependents'):
                    self.main_controller.refresh_dependents()
                QMessageBox.information(self.main_controller, "Thành công", f"Đã xóa thành công bản ghi {pk_val}.")
            except Exception as e:
                QMessageBox.critical(self.main_controller, "Lỗi", f"Không thể xóa bản ghi: {e}")

    def action_luu(self):
        try:
            conn = sqlite3.connect(DB_PATH)
            cursor = conn.cursor()
            
            if self.db_table == "LIVESTREAM_SAN_PHAM":
                ma_live = get_widget_value(self.ui.cbLivestream, "combo_text")
                ma_sp = get_widget_value(self.ui.cbSanPham, "combo_text")
                if not ma_live or not ma_sp:
                    QMessageBox.warning(self.main_controller, "Lỗi dữ liệu", "Vui lòng chọn livestream và sản phẩm!")
                    conn.close()
                    return
                
                if self.current_action == "THEM":
                    cursor.execute("SELECT 1 FROM LIVESTREAM_SAN_PHAM WHERE MaLive = ? AND MaSP = ?", (ma_live, ma_sp))
                    if cursor.fetchone():
                        QMessageBox.warning(self.main_controller, "Lỗi", "Bản ghi liên kết này đã tồn tại!")
                        conn.close()
                        return
                    cursor.execute("INSERT INTO LIVESTREAM_SAN_PHAM (MaLive, MaSP) VALUES (?, ?)", (ma_live, ma_sp))
                elif self.current_action == "SUA":
                    parts = self.selected_pk.split("-")
                    cursor.execute("DELETE FROM LIVESTREAM_SAN_PHAM WHERE MaLive = ? AND MaSP = ?", (parts[0], parts[1]))
                    cursor.execute("INSERT INTO LIVESTREAM_SAN_PHAM (MaLive, MaSP) VALUES (?, ?)", (ma_live, ma_sp))
                
                conn.commit()
                conn.close()
                self.current_action = None
                self.set_form_enabled(False)
                self.refresh_callback()
                if hasattr(self.main_controller, 'refresh_dependents'):
                    self.main_controller.refresh_dependents()
                QMessageBox.information(self.main_controller, "Thành công", "Lưu liên kết livestream - sản phẩm thành công.")
                return

            elif self.db_table == "CHI_TIET_DON_HANG":
                ma_dh = get_widget_value(self.ui.cbDonHang, "combo_text")
                ma_sp = get_widget_value(self.ui.cbSanPham, "combo_text")
                so_luong = get_widget_value(self.ui.spinSoLuong, "spin")
                don_gia_str = self.ui.txtGiaBan.text().strip().replace(",", "")
                
                if not ma_dh or not ma_sp:
                    QMessageBox.warning(self.main_controller, "Lỗi dữ liệu", "Vui lòng chọn đơn hàng và sản phẩm!")
                    conn.close()
                    return
                try:
                    don_gia = float(don_gia_str) if don_gia_str else 0.0
                except ValueError:
                    QMessageBox.warning(self.main_controller, "Lỗi dữ liệu", "Giá bán phải là số hợp lệ!")
                    conn.close()
                    return
                
                thanh_tien = so_luong * don_gia
                
                if self.current_action == "THEM":
                    cursor.execute("SELECT 1 FROM CHI_TIET_DON_HANG WHERE MaDonHang = ? AND MaSP = ?", (ma_dh, ma_sp))
                    if cursor.fetchone():
                        QMessageBox.warning(self.main_controller, "Lỗi", "Chi tiết đơn hàng này đã tồn tại!")
                        conn.close()
                        return
                    cursor.execute("INSERT INTO CHI_TIET_DON_HANG (MaDonHang, MaSP, SoLuong, DonGia, ThanhTien) VALUES (?, ?, ?, ?, ?)",
                                   (ma_dh, ma_sp, so_luong, don_gia, thanh_tien))
                elif self.current_action == "SUA":
                    parts = self.selected_pk.split("-")
                    cursor.execute("DELETE FROM CHI_TIET_DON_HANG WHERE MaDonHang = ? AND MaSP = ?", (parts[0], parts[1]))
                    cursor.execute("INSERT INTO CHI_TIET_DON_HANG (MaDonHang, MaSP, SoLuong, DonGia, ThanhTien) VALUES (?, ?, ?, ?, ?)",
                                   (ma_dh, ma_sp, so_luong, don_gia, thanh_tien))
                
                # Update Totals on DON_HANG and HOA_DON (đã trừ voucher nếu có)
                recalc_order_total(cursor, ma_dh)
                
                conn.commit()
                conn.close()
                self.current_action = None
                self.set_form_enabled(False)
                self.refresh_callback()
                if hasattr(self.main_controller, 'refresh_dependents'):
                    self.main_controller.refresh_dependents()
                QMessageBox.information(self.main_controller, "Thành công", "Lưu chi tiết đơn hàng thành công.")
                return

            elif self.db_table == "DON_HANG":
                ten_kh = self.ui.txtTenKhachHang.text().strip()
                sdt = self.ui.txtSoDienThoai.text().strip()
                dia_chi = self.ui.txtDiaChiGiao.text().strip()
                ngay_dat = self.ui.dateNgayDat.dateTime().toString("yyyy-MM-dd HH:mm:ss")
                ma_voucher = get_widget_value(self.ui.cbVoucher, "combo_text")
                trang_thai = self.ui.cbTrangThai.currentText()
                pt_tt = self.ui.cbPaymentMode.currentText()
                
                if not ten_kh or not sdt:
                    QMessageBox.warning(self.main_controller, "Lỗi dữ liệu", "Vui lòng nhập tên và số điện thoại khách hàng!")
                    conn.close()
                    return
                
                # Check / Insert Customer
                cursor.execute("SELECT MaKhachHang FROM KHACH_HANG WHERE SoDienThoai = ?", (sdt,))
                cust_res = cursor.fetchone()
                if cust_res:
                    ma_kh = cust_res[0]
                    cursor.execute("UPDATE KHACH_HANG SET HoTen = ?, DiaChi = ? WHERE MaKhachHang = ?", (ten_kh, dia_chi, ma_kh))
                else:
                    ma_kh = self.generate_new_pk_generic("KHACH_HANG", "MaKhachHang", "KH")
                    cursor.execute("INSERT INTO KHACH_HANG (MaKhachHang, HoTen, SoDienThoai, Email, DiaChi) VALUES (?, ?, ?, ?, ?)",
                                   (ma_kh, ten_kh, sdt, "", dia_chi))
                
                # Tổng tiền nhập tay (chi tiết đơn hàng sẽ ghi đè khi đơn có sản phẩm)
                tong_tien_str = self.ui.txtTongTien.text().strip().replace(",", "") if hasattr(self.ui, 'txtTongTien') else ""
                try:
                    tong_tien_manual = float(tong_tien_str) if tong_tien_str else None
                except ValueError:
                    tong_tien_manual = None

                old_status = None
                if self.current_action == "THEM":
                    ma_dh = self.generate_new_pk_generic("DON_HANG", "MaDonHang", "DH")
                    cursor.execute("INSERT INTO DON_HANG (MaDonHang, MaKhachHang, MaVoucher, NgayDat, TongTien, TrangThaiDH) VALUES (?, ?, ?, ?, ?, ?)",
                                   (ma_dh, ma_kh, ma_voucher, ngay_dat, tong_tien_manual or 0.0, trang_thai))

                    # Create corresponding Invoice
                    ma_hd = self.generate_new_pk_generic("HOA_DON", "MaHoaDon", "HD")
                    cursor.execute("INSERT INTO HOA_DON (MaHoaDon, MaDonHang, PhuongThucTT, TongTien, ThoiGianLap, TrangThaiHD) VALUES (?, ?, ?, ?, ?, 'Chưa thanh toán')",
                                   (ma_hd, ma_dh, pt_tt, tong_tien_manual or 0.0, ngay_dat))
                elif self.current_action == "SUA":
                    ma_dh = self.selected_pk
                    cursor.execute("SELECT TrangThaiDH FROM DON_HANG WHERE MaDonHang = ?", (ma_dh,))
                    res = cursor.fetchone()
                    old_status = res[0] if res else None
                    cursor.execute("UPDATE DON_HANG SET MaKhachHang = ?, MaVoucher = ?, NgayDat = ?, TrangThaiDH = ? WHERE MaDonHang = ?",
                                   (ma_kh, ma_voucher, ngay_dat, trang_thai, ma_dh))
                    if tong_tien_manual is not None:
                        cursor.execute("UPDATE DON_HANG SET TongTien = ? WHERE MaDonHang = ?", (tong_tien_manual, ma_dh))
                    cursor.execute("UPDATE HOA_DON SET PhuongThucTT = ? WHERE MaDonHang = ?", (pt_tt, ma_dh))

                # Đồng bộ tồn kho + trạng thái hóa đơn theo trạng thái đơn
                apply_order_status_effects(cursor, ma_dh, old_status, trang_thai)
                if trang_thai not in ("Đã giao", "Đã hủy"):
                    cursor.execute("UPDATE HOA_DON SET TrangThaiHD = 'Chưa thanh toán' WHERE MaDonHang = ?", (ma_dh,))

                # Đơn có chi tiết thì tổng tiền tính từ chi tiết (trừ voucher)
                cursor.execute("SELECT 1 FROM CHI_TIET_DON_HANG WHERE MaDonHang = ? LIMIT 1", (ma_dh,))
                if cursor.fetchone():
                    recalc_order_total(cursor, ma_dh)

                conn.commit()
                conn.close()
                self.current_action = None
                self.set_form_enabled(False)
                self.refresh_callback()
                if hasattr(self.main_controller, 'refresh_dependents'):
                    self.main_controller.refresh_dependents()
                QMessageBox.information(self.main_controller, "Thành công", "Lưu thông tin đơn hàng thành công.")
                return

            else:
                # Standard save flow
                vals = {}
                for db_col, widget, widget_type in self.col_mappings:
                    if widget:
                        vals[db_col] = get_widget_value(widget, widget_type)
                
                # Combo voucher bên Đơn hàng hiển thị theo TenVoucher — mặc định dùng mã
                if self.db_table == "VOUCHER" and not vals.get("TenVoucher"):
                    vals["TenVoucher"] = vals.get("MaVoucher") or self.selected_pk

                pk_widget = self.get_pk_widget()
                if self.current_action == "THEM":
                    if pk_widget:
                        pk_val = get_widget_value(pk_widget, self.get_pk_widget_type())
                        if not pk_val:
                            QMessageBox.warning(self.main_controller, "Lỗi dữ liệu", "Vui lòng nhập Mã khóa chính!")
                            conn.close()
                            return
                        vals[self.pk_col] = pk_val
                    else:
                        pk_val = self.generate_new_pk()
                        vals[self.pk_col] = pk_val
                    
                    cols = list(vals.keys())
                    placeholders = [f":" + c for c in cols]
                    query = f"INSERT INTO {self.db_table} ({', '.join(cols)}) VALUES ({', '.join(placeholders)})"
                    cursor.execute(query, vals)
                    
                elif self.current_action == "SUA":
                    if pk_widget:
                        pk_val = get_widget_value(pk_widget, self.get_pk_widget_type())
                    else:
                        pk_val = self.selected_pk
                        
                    if not pk_val:
                        QMessageBox.warning(self.main_controller, "Lỗi", "Không xác định được Mã bản ghi cần cập nhật!")
                        conn.close()
                        return
                        
                    set_clause = ", ".join([f"{col} = :{col}" for col in vals.keys() if col != self.pk_col])
                    vals[self.pk_col] = pk_val
                    query = f"UPDATE {self.db_table} SET {set_clause} WHERE {self.pk_col} = :{self.pk_col}"
                    cursor.execute(query, vals)
                
                conn.commit()
                conn.close()
                self.current_action = None
                self.set_form_enabled(False)
                self.refresh_callback()
                if hasattr(self.main_controller, 'refresh_dependents'):
                    self.main_controller.refresh_dependents()
                QMessageBox.information(self.main_controller, "Thành công", "Lưu dữ liệu thành công.")
        except Exception as e:
            QMessageBox.critical(self.main_controller, "Lỗi", f"Không thể lưu dữ liệu: {e}")

    def action_huy(self):
        self.current_action = None
        self.clear_form()
        self.set_form_enabled(False)
        self.selected_pk = None
        # Nạp lại bảng thay vì xóa trắng để không gây cảm giác mất dữ liệu
        self.refresh_callback()
        if hasattr(self.ui, "lblStatus"):
            self.ui.lblStatus.setText("Trạng thái: Đã hủy thao tác.")

    def action_refresh(self):
        if hasattr(self.ui, 'txtSearch'):
            self.ui.txtSearch.clear()
        self.refresh_callback()

    def action_search(self):
        if not hasattr(self.ui, 'txtSearch'):
            return
        search_text = self.ui.txtSearch.text().strip()
        if not search_text:
            self.refresh_callback()
            return
        
        try:
            conn = sqlite3.connect(DB_PATH)
            cursor = conn.cursor()
            
            if self.db_table == "BINH_LUAN":
                query = """
                    SELECT MaBinhLuan, MaLive, NguoiBinhLuan, NoiDung, ThoiGian 
                    FROM BINH_LUAN 
                    WHERE NguoiBinhLuan LIKE ? OR NoiDung LIKE ? OR MaLive LIKE ?
                """
                cursor.execute(query, (f"%{search_text}%", f"%{search_text}%", f"%{search_text}%"))
            elif self.db_table == "KHACH_HANG":
                query = """
                    SELECT MaKhachHang, HoTen, SoDienThoai, Email, DiaChi, '2026-07-15' 
                    FROM KHACH_HANG 
                    WHERE HoTen LIKE ? OR SoDienThoai LIKE ? OR Email LIKE ? OR DiaChi LIKE ?
                """
                cursor.execute(query, (f"%{search_text}%", f"%{search_text}%", f"%{search_text}%", f"%{search_text}%"))
            elif self.db_table == "LIVESTREAM":
                query = """
                    SELECT MaLive, TenLive, ThoiGianBatDau, MaNguoiBan, TrangThai 
                    FROM LIVESTREAM 
                    WHERE TenLive LIKE ? OR MaLive LIKE ? OR MaNguoiBan LIKE ? OR TrangThai LIKE ?
                """
                cursor.execute(query, (f"%{search_text}%", f"%{search_text}%", f"%{search_text}%", f"%{search_text}%"))
            elif self.db_table == "LIVESTREAM_SAN_PHAM":
                query = """
                    SELECT 
                        (lsp.MaLive || '-' || lsp.MaSP) AS MaChiTietLive,
                        l.TenLive,
                        lsp.MaSP,
                        p.TenSP,
                        p.SoLuongTon,
                        p.GiaBan
                    FROM LIVESTREAM_SAN_PHAM lsp
                    JOIN LIVESTREAM l ON lsp.MaLive = l.MaLive
                    JOIN SAN_PHAM p ON lsp.MaSP = p.MaSP
                    WHERE l.TenLive LIKE ? OR p.TenSP LIKE ? OR lsp.MaLive LIKE ? OR lsp.MaSP LIKE ?
                """
                cursor.execute(query, (f"%{search_text}%", f"%{search_text}%", f"%{search_text}%", f"%{search_text}%"))
            elif self.db_table == "DON_HANG":
                query = """
                    SELECT 
                        dh.MaDonHang,
                        dh.NgayDat,
                        kh.HoTen,
                        kh.SoDienThoai,
                        dh.TongTien,
                        COALESCE(v.TenVoucher, dh.MaVoucher),
                        dh.TrangThaiDH,
                        COALESCE((SELECT GROUP_CONCAT(p.TenSP, ', ')
                                  FROM CHI_TIET_DON_HANG ct JOIN SAN_PHAM p ON ct.MaSP = p.MaSP
                                  WHERE ct.MaDonHang = dh.MaDonHang), '(chưa có SP)')
                    FROM DON_HANG dh
                    LEFT JOIN KHACH_HANG kh ON dh.MaKhachHang = kh.MaKhachHang
                    LEFT JOIN VOUCHER v ON dh.MaVoucher = v.MaVoucher
                    WHERE dh.MaDonHang LIKE ? OR kh.HoTen LIKE ? OR kh.SoDienThoai LIKE ? OR dh.TrangThaiDH LIKE ?
                """
                cursor.execute(query, (f"%{search_text}%", f"%{search_text}%", f"%{search_text}%", f"%{search_text}%"))
            elif self.db_table == "CHI_TIET_DON_HANG":
                query = """
                    SELECT 
                        (ct.MaDonHang || '-' || ct.MaSP) AS MaChiTietDH,
                        ct.MaDonHang,
                        ct.MaSP,
                        p.TenSP,
                        ct.SoLuong,
                        ct.DonGia,
                        ct.ThanhTien,
                        COALESCE(v.TenVoucher, dh.MaVoucher, 'Không')
                    FROM CHI_TIET_DON_HANG ct
                    JOIN SAN_PHAM p ON ct.MaSP = p.MaSP
                    LEFT JOIN DON_HANG dh ON ct.MaDonHang = dh.MaDonHang
                    LEFT JOIN VOUCHER v ON dh.MaVoucher = v.MaVoucher
                    WHERE ct.MaDonHang LIKE ? OR ct.MaSP LIKE ? OR p.TenSP LIKE ?
                """
                cursor.execute(query, (f"%{search_text}%", f"%{search_text}%", f"%{search_text}%"))
            elif self.db_table == "HOA_DON":
                query = """
                    SELECT MaHoaDon, MaDonHang, ThoiGianLap, TongTien, PhuongThucTT, TrangThaiHD 
                    FROM HOA_DON 
                    WHERE MaHoaDon LIKE ? OR MaDonHang LIKE ? OR PhuongThucTT LIKE ? OR TrangThaiHD LIKE ?
                """
                cursor.execute(query, (f"%{search_text}%", f"%{search_text}%", f"%{search_text}%", f"%{search_text}%"))
            elif self.db_table == "VOUCHER":
                query = """
                    SELECT
                        MaVoucher,
                        COALESCE(LoaiUuDai, 'Số tiền cố định (VND)') AS LoaiUuDai,
                        GiaTriGiam,
                        DieuKienApDung,
                        COALESCE(GiamToiDa, GiaTriGiam) AS GiamToiDa,
                        NgayBatDau,
                        NgayKetThuc,
                        TrangThai
                    FROM VOUCHER
                    WHERE MaVoucher LIKE ? OR TenVoucher LIKE ? OR TrangThai LIKE ?
                """
                cursor.execute(query, (f"%{search_text}%", f"%{search_text}%", f"%{search_text}%"))
            else:
                conn.close()
                self.main_controller.search_table_generic(self.table_widget, self.db_table, self.pk_col, search_text, self.refresh_callback)
                return
            
            results = cursor.fetchall()
            conn.close()
            self.main_controller.populate_table(self.table_widget, results)
            if hasattr(self.ui, 'lblStatus'):
                self.ui.lblStatus.setText(f"Trạng thái: Tìm thấy {len(results)} kết quả.")
        except Exception as e:
            print(f"Error action_search for {self.db_table}: {e}")

    def action_xuat_excel(self):
        path, _ = QtWidgets.QFileDialog.getSaveFileName(
            self.main_controller, "Lưu file Excel", "", "Excel Files (*.xlsx);;All Files (*)"
        )
        if not path:
            return

        try:
            from openpyxl import Workbook
            wb = Workbook()
            ws = wb.active
            ws.title = self.db_table[:30]

            headers = []
            for col in range(self.table_widget.columnCount()):
                header_item = self.table_widget.horizontalHeaderItem(col)
                headers.append(header_item.text() if header_item else f"Cột {col+1}")
            ws.append(headers)

            for row in range(self.table_widget.rowCount()):
                row_data = []
                for col in range(self.table_widget.columnCount()):
                    item = self.table_widget.item(row, col)
                    row_data.append(item.text() if item else "")
                ws.append(row_data)

            wb.save(path)
            QMessageBox.information(self.main_controller, "Thành công", f"Đã xuất dữ liệu Excel thành công tại:\n{path}")
        except ImportError:
            QMessageBox.critical(self.main_controller, "Lỗi", "Vui lòng cài đặt openpyxl bằng lệnh: pip install openpyxl")
        except Exception as e:
            QMessageBox.critical(self.main_controller, "Lỗi", f"Không thể lưu file Excel: {e}")


